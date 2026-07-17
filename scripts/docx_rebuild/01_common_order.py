from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import copy, os, re, zipfile, hashlib, json, shutil, textwrap, math

repo_root = Path.cwd()
outdir = repo_root / '.tmp-official-docx'
outdir.mkdir(exist_ok=True)

def clear_paragraph(p):
    pPr = p._p.pPr
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)

def set_run_font(run, name="Arial", size=9, bold=None, italic=None, color=None):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('ascii','hAnsi','eastAsia','cs'):
        rFonts.set(qn(f'w:{attr}'), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = color

def set_para(p, segments, align=None, before=0, after=0, line=1.0, keep=False):
    clear_paragraph(p)
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    for seg in segments:
        if isinstance(seg, str):
            seg = {"text": seg}
        r = p.add_run(seg.get("text",""))
        set_run_font(r, seg.get("font","Arial"), seg.get("size",9), seg.get("bold"), seg.get("italic"), seg.get("color"))
    return p

def ensure_paragraphs(cell, count):
    while len(cell.paragraphs) < count:
        cell.add_paragraph()
    return cell.paragraphs

def clear_cell(cell, keep_first=True):
    tc = cell._tc
    for child in list(tc):
        if child.tag == qn('w:tcPr'):
            continue
        tc.remove(child)
    p = OxmlElement('w:p')
    tc.append(p)
    return cell.paragraphs[0]

def set_cell_margins(cell, top=40, start=50, bottom=40, end=50):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in (('top',top),('start',start),('bottom',bottom),('end',end)):
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if edge in edges:
            edge_data=edges[edge]
            tag = qn(f'w:{edge}')
            element = tcBorders.find(tag)
            if element is None:
                element = OxmlElement(f'w:{edge}')
                tcBorders.append(element)
            for key,val in edge_data.items():
                element.set(qn(f'w:{key}'), str(val))

def set_table_borders(table, val='single', sz=4, color='000000'):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={'val':val,'sz':sz,'color':color}, bottom={'val':val,'sz':sz,'color':color}, left={'val':val,'sz':sz,'color':color}, right={'val':val,'sz':sz,'color':color})

def set_no_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'},insideH={'val':'nil'},insideV={'val':'nil'})

def set_row_height(row, cm, exact=False):
    row.height = Cm(cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST

def set_doc_defaults(doc, font="Arial", size=9):
    normal = doc.styles['Normal']
    normal.font.name = font
    normal.font.size = Pt(size)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), font)

def sanitize_props(doc, title):
    cp=doc.core_properties
    cp.author="CONSOLARE"
    cp.last_modified_by="CONSOLARE"
    cp.title=title
    cp.subject="Modelo oficial"
    cp.comments=""
    cp.keywords=""

def style_all_table_runs(doc, default_font="Arial", min_size=9):
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    if p.paragraph_format.line_spacing is None:
                        p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        if r.text:
                            name = r.font.name or default_font
                            size = r.font.size.pt if r.font.size else min_size
                            if size < min_size and size > 2:
                                size = min_size
                            set_run_font(r, name, size, r.bold, r.italic)

def remove_all_body_except_sectpr(doc):
    body = doc._element.body
    sectPr = body.sectPr
    for child in list(body):
        if child is not sectPr:
            body.remove(child)

def add_text(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font="Arial", italic=False):
    p = cell.paragraphs[0]
    set_para(p, [{"text": text, "size": size, "bold": bold, "font": font, "italic": italic}], align=align)
    return p

order_src = repo_root / 'public/templates/official/sepultamento/ordem-sepultamento.docx'
def make_order(kind):
    doc = Document(order_src)
    set_doc_defaults(doc, "Arial", 9)
    t = doc.tables[0]
    c=t.cell(1,0)
    ps=ensure_paragraphs(c,3)
    set_para(ps[0],[{"text":"Concessão:","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(ps[1],[],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(ps[2],[{"text":"□ SIM      ■ NÃO" if kind=="quadra" else "■ SIM      □ NÃO","size":10}],align=WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(1,2); ps=ensure_paragraphs(c,3)
    set_para(ps[0],[{"text":"Quadra geral c/gaveta:","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(ps[1],[],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(ps[2],[{"text":"■ SIM      □ NÃO" if kind=="quadra" else "□ SIM      ■ NÃO","size":10}],align=WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(1,6); p=clear_cell(c)
    set_para(p,[{"text":"Sala:","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(c.add_paragraph(),[{"text":"{salaVelorio}","size":10,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell,label,ph in [(t.cell(2,0),"N° de Inscrição (GSCEMI):","{inscrGS}"),(t.cell(2,2),"Nº da DO:","{numDO}"),(t.cell(2,6),"Nº Placa de Identificação:","{placa}")]:
        p=clear_cell(cell)
        set_para(p,[{"text":label,"size":8.5}],align=WD_ALIGN_PARAGRAPH.CENTER)
        set_para(cell.add_paragraph(),[{"text":ph,"size":10,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(3,0); p=clear_cell(c)
    set_para(p,[{"text":"NOME DA PESSOA FALECIDA:","size":9,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(c.add_paragraph(),[{"text":"{nomeFal}","size":14,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell,label,ph,sz in [(t.cell(4,0),"DATA DO SEPULTAMENTO:","{dataSep}",11),(t.cell(4,2),"HORA DO SEPULTAMENTO:","{horaSep}",10)]:
        p=clear_cell(cell)
        set_para(p,[{"text":label,"size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
        set_para(cell.add_paragraph(),[{"text":ph,"size":sz,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(4,8); p=clear_cell(c)
    set_para(p,[{"text":"Família Presente:","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(c.add_paragraph(),[{"text":"□ Sim      □ Não","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(t.cell(5,0).paragraphs[0],[{"text":"AUTORIZAÇÃO DE SEPULTAMENTO","size":10,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(6,0); clear_cell(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_para(c.paragraphs[0],[{"text":"Eu, ","size":9},{"text":"{nomeResp}","size":9,"bold":True},{"text":" CPF: ","size":9},{"text":"{cpfResp}","size":9,"bold":True},{"text":" Profissão: ____________________","size":9}], after=3)
    set_para(c.add_paragraph(),[{"text":"Endereço: ","size":9},{"text":"{endResp}","size":9,"bold":True}],after=2)
    set_para(c.add_paragraph(),[{"text":"Telefone: ","size":9},{"text":"{telResp}","size":9,"bold":True}],after=5)
    set_para(c.add_paragraph(),[{"text":"Na qualidade de ","size":9},{"text":"{parent}","size":9,"bold":True},{"text":" autorizo o sepultamento junto à administração do Cemitério Santana, nesta data.","size":9}],after=4)
    set_para(c.add_paragraph(),[{"text":"Isento a empresa ","size":9,"italic":True},{"text":"Consolare ","size":9,"bold":True,"italic":True},{"text":"de qualquer responsabilidade e assumo total responsabilidade pelos serviços realizados de exumação, corte de gaveta, remoção da alça, rompimento do portão de sepultamentos e quaisquer outros serviços que se fizerem necessários.","size":9,"italic":True}],after=4)
    set_para(c.add_paragraph(),[{"text":"Por ser de minha livre e espontânea vontade, firmo a presente autorização.","size":9,"italic":True}],after=8)
    set_para(c.add_paragraph(),[{"text":"Assinatura do responsável pela autorização: ________________________________________________","size":9}],after=8)
    set_para(c.add_paragraph(),[{"text":"São Paulo, {dataExt}.","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
    set_para(c.add_paragraph(),[{"text":"Contratação:  □ Social    □ Doador    □ Padrão    □ Popular    □ Luxo    □ De Fora","size":8.5}],after=3)
    set_para(c.add_paragraph(),[{"text":"Obs.: ___________________________    Nota de Contratação nº: ","size":8.5},{"text":"{nota}","size":8.5,"bold":True}],after=3)
    set_para(c.add_paragraph(),[{"text":"COVID/Lacrado:  □ sim    □ não      Empresa / Bloco / Agência: ","size":8.5},{"text":"{funeraria}","size":8.5,"bold":True}])
    for cell,label,ph in [(t.cell(8,0),"N° quadra / Rua","{quadraRua}"),(t.cell(8,1),"Nº Sepultura",""),(t.cell(8,3),"Rua",""),(t.cell(8,4),"Nº gaveta","{gaveta}")]:
        p=clear_cell(cell)
        set_para(p,[{"text":label,"size":8.5}],align=WD_ALIGN_PARAGRAPH.CENTER)
        set_para(cell.add_paragraph(),[{"text":ph,"size":10.5,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(8,7); clear_cell(c)
    nt=c.add_table(rows=3, cols=2); nt.autofit=False; set_no_borders(nt)
    for i,(a,b) in enumerate([("Óbito","Concessão"),("L: {livroObito}","L: "),("F: ","F: ")]):
        add_text(nt.cell(i,0),a,size=8.5,bold=True)
        add_text(nt.cell(i,1),b,size=8.5,bold=True)
    set_para(t.cell(9,0).paragraphs[0],[{"text":"RESPONSÁVEIS PELO SEPULTAMENTO:","size":9,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell,text in [(t.cell(10,0),"Resp. Sepultamento:"),(t.cell(10,2),"visto sepultador"),(t.cell(10,5),"visto sepultador"),(t.cell(10,9),"visto sepultador")]:
        set_para(clear_cell(cell),[{"text":text,"size":8.5}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(t.cell(11,0).paragraphs[0],[{"text":"ADMINISTRAÇÃO","size":9,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(clear_cell(t.cell(12,0)),[{"text":"Assinatura da Administração: __________________________    Data: ","size":8.5},{"text":"{dataAtual}","size":8.5,"bold":True}],align=WD_ALIGN_PARAGRAPH.CENTER)
    style_all_table_runs(doc,"Arial",8.5)
    set_run_font(t.cell(0,0).paragraphs[0].runs[0],"Arial",12,True)
    for idx,h in {0:0.45,1:1.0,2:0.75,3:0.75,4:0.85,5:0.4,7:0.25,8:1.25,9:0.38,10:0.62,11:0.38,12:0.65}.items():
        set_row_height(t.rows[idx],h,exact=False)
    sanitize_props(doc, f"Ordem de Sepultamento - {'Quadra Geral' if kind=='quadra' else 'Jazigo'}")
    out=outdir/f"ordem-sepultamento{'-jazigo' if kind=='jazigo' else ''}.docx"
    doc.save(out)

make_order("quadra")
make_order("jazigo")
