from docx.enum.section import WD_SECTION

oss_src=repo_root / 'public/templates/official/ossuario/aquisicao-renovacao-ossuario.docx'

def set_cell_text(cell, label, value=None, label_size=8.5, value_size=11, font="Arial", value_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER):
    clear_cell(cell)
    set_para(cell.paragraphs[0],[{"text":label,"font":font,"size":label_size,"bold":True}],align=align)
    if value is not None:
        set_para(cell.add_paragraph(),[{"text":value,"font":font,"size":value_size,"bold":value_bold}],align=align)

def make_ossuario(kind):
    doc=Document(oss_src)
    set_doc_defaults(doc,"Arial",10)
    remove_all_body_except_sectpr(doc)
    section=doc.sections[0]
    section.top_margin=Cm(2.5)
    section.bottom_margin=Cm(1.8)
    section.left_margin=Cm(1.5)
    section.right_margin=Cm(1.5)
    outer=doc.add_table(rows=1, cols=1)
    outer.autofit=False
    outer.columns[0].width=Cm(17.7)
    set_table_borders(outer,'single',6,'666666')
    cell=outer.cell(0,0)
    set_cell_margins(cell,180,180,120,180)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    clear_cell(cell)
    op=cell.add_table(rows=1,cols=2)
    op.autofit=False
    set_no_borders(op)
    left="■ 1º ALUGUEL" if kind=="aquisicao" else "□ 1º ALUGUEL"
    right="□ RENOVAÇÃO" if kind=="aquisicao" else "■ RENOVAÇÃO"
    add_text(op.cell(0,0),left,10,True,WD_ALIGN_PARAGRAPH.CENTER)
    add_text(op.cell(0,1),right,10,True,WD_ALIGN_PARAGRAPH.CENTER)
    details=[
        ([{"text":"Concessionário: ","size":10},{"text":"{nomeConc}","size":10.5,"bold":True}],6),
        ([{"text":"CPF: ","size":10},{"text":"{cpfConc}","size":10.5,"bold":True}],4),
        ([{"text":"Endereço: ","size":10},{"text":"{endConc}","size":10.5,"bold":True}],4),
        ([{"text":"Telefone: ","size":10},{"text":"{telConc}","size":10.5,"bold":True}],12),
        ([{"text":"Despojos de: ","size":10},{"text":"{nomeFal}","size":10.5,"bold":True}],16),
    ]
    for segments,after in details:
        set_para(cell.add_paragraph(),segments,after=after)
    phrase="AQUISIÇÃO de ossuário determinado" if kind=="aquisicao" else "RENOVAÇÃO de ossuário determinado"
    set_para(cell.add_paragraph(),[{"text":phrase,"size":10},{"text":" no valor vigente.","size":10}],align=WD_ALIGN_PARAGRAPH.CENTER,after=10)
    tbl=cell.add_table(rows=3,cols=6)
    tbl.autofit=False
    set_table_borders(tbl,'single',4,'000000')
    title=tbl.cell(0,0)
    for c2 in tbl.rows[0].cells[1:]:
        title=title.merge(c2)
    set_cell_text(tbl.cell(0,0),"OSSUÁRIO",None,9,9)
    tbl.cell(1,1).merge(tbl.cell(1,2))
    tbl.cell(1,4).merge(tbl.cell(1,5))
    date_label="AQUISIÇÃO EM:" if kind=="aquisicao" else "RENOVAÇÃO EM:"
    set_cell_text(tbl.cell(1,0),"BLOCO/GALERIA","{bloco}",8,12)
    set_cell_text(tbl.cell(1,1),"NÚMERO","{numeroOssuario}",8,12)
    set_cell_text(tbl.cell(1,3),date_label,"{dataAquisicao}",8,11)
    set_cell_text(tbl.cell(1,4),"VENCIMENTO EM:","{dataVencimento}",8,11)
    tbl.cell(2,0).merge(tbl.cell(2,1))
    tbl.cell(2,2).merge(tbl.cell(2,3))
    set_cell_text(tbl.cell(2,0),"Nº INSCRIÇÃO (GSCEMI)","{inscrGS}",7.5,9.5)
    set_cell_text(tbl.cell(2,2),"Nº PLACA DE IDENTIFICAÇÃO","{placa}",7.5,9.5)
    set_cell_text(tbl.cell(2,4),"Nº DO LIVRO","{livro}",7.5,9.5)
    set_cell_text(tbl.cell(2,5),"FOLHA","{folha}",7.5,9.5)
    for row in tbl.rows:
        for c in row.cells:
            set_cell_margins(c,40,35,40,35)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_row_height(tbl.rows[0],0.38,True)
    set_row_height(tbl.rows[1],1.25,False)
    set_row_height(tbl.rows[2],0.95,False)
    set_para(cell.add_paragraph(),[{"text":"Obs.: ______________________________________________________________________________","size":9}],before=8,after=12)
    set_para(cell.add_paragraph(),[{"text":"São Paulo, {dataExt}.","size":9.5}],align=WD_ALIGN_PARAGRAPH.CENTER,after=14)
    set_para(cell.add_paragraph(),[{"text":"__________________________________","size":9}],align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(cell.add_paragraph(),[{"text":"Assinatura do Funcionário","size":8}],align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
    set_para(cell.add_paragraph(),[{"text":"................................................................................................................................................................................","size":7}],after=6)
    receipt=cell.add_table(rows=3,cols=6)
    receipt.autofit=False
    set_table_borders(receipt,'single',4,'000000')
    title=receipt.cell(0,0)
    for c2 in receipt.rows[0].cells[1:]:
        title=title.merge(c2)
    set_cell_text(receipt.cell(0,0),"OSSUÁRIO - VIA DE CONTROLE",None,8.5,8.5)
    receipt.cell(1,1).merge(receipt.cell(1,2))
    receipt.cell(1,4).merge(receipt.cell(1,5))
    set_cell_text(receipt.cell(1,0),"BLOCO/GALERIA","{bloco}",7.5,9.5)
    set_cell_text(receipt.cell(1,1),"NÚMERO","{numeroOssuario}",7.5,9.5)
    set_cell_text(receipt.cell(1,3),date_label,"{dataAquisicao}",7.5,9)
    set_cell_text(receipt.cell(1,4),"VENCIMENTO EM:","{dataVencimento}",7.5,9)
    receipt.cell(2,0).merge(receipt.cell(2,1))
    receipt.cell(2,2).merge(receipt.cell(2,3))
    set_cell_text(receipt.cell(2,0),"Nº INSCRIÇÃO (GSCEMI)","{inscrGS}",7,8.5)
    set_cell_text(receipt.cell(2,2),"Nº PLACA","{placa}",7,8.5)
    set_cell_text(receipt.cell(2,4),"LIVRO","{livro}",7,8.5)
    set_cell_text(receipt.cell(2,5),"FOLHA","{folha}",7,8.5)
    for row in receipt.rows:
        for c in row.cells:
            set_cell_margins(c,25,25,25,25)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_row_height(receipt.rows[0],0.35,True)
    set_row_height(receipt.rows[1],0.85,False)
    set_row_height(receipt.rows[2],0.7,False)
    sanitize_props(doc,"Ossuário - Aquisição" if kind=="aquisicao" else "Ossuário - Renovação")
    filename="aquisicao-renovacao-ossuario.docx" if kind=="aquisicao" else "renovacao-ossuario.docx"
    doc.save(outdir/filename)

make_ossuario('aquisicao')
make_ossuario('renovacao')
